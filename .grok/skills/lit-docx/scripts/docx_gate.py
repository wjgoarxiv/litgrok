#!/usr/bin/env python3
"""One pass/fail gate for a converted .docx and the Markdown it came from.

It reads five things and prints one JSON report (exit 0 = PASS, 1 = FAIL):

  package    the file is a readable ZIP with a document part, no NaN attribute values, and
             python-docx reopens it
  content    the body has text, every Markdown heading reached the document, no unfilled blank
             ([제품명], XXX, TBD, ○○, {{ … }}) is left, and no frontmatter was printed as text
  lint       slop_lint's prose rules on the source (journal outline rules only for a manuscript)
             and its DOCX audits: numeric-column alignment (OF-302), Hangul font pairing, and on
             a publisher profile that profile's design rules
  output     docx_layout's checks: heading labels and order, the notice's colon, the restraint rules
             of a tonality document, and with --layout the page checks on a LibreOffice render (fill,
             spill onto a near-empty page, stranded headings, split short lists, a heading apart from
             its table, table and figure splits, component variety, cover)
  compare    with --compare OTHER.docx, two tonalities of one source must differ in at least three
             structural features

Usage:
  node .grok/skills/lit-docx/scripts/run.mjs docx_gate.py report.docx --source report.md --layout
  node .grok/skills/lit-docx/scripts/run.mjs docx_gate.py paper.docx --source paper.md --publisher elsevier --kind manuscript

A PASS means no defect was found. Whether the pages read well is decided by looking at them.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document

import docx_layout
import slop_lint
from convert_md_to_docx import parse_frontmatter

SKILL = Path(__file__).resolve().parents[1]
HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
# A bracket holding a digit is a citation marker ([1], [S3]); a blank holds none.
BLANK = re.compile(r"\[(?![^\]\n]*\d)[^\]\n]{1,24}\]|\bX{3,}\b|\bTBD\b|\bTODO\b|○○|OOO|\{\{[^}]*\}\}|_{4,}|＿{2,}")
LEAK = re.compile(r"^(?:---|title:|subtitle:|author:|authors:|date:|tonality:)\s*", re.M)
# Audits every profile shares; the rest belong to a named publisher profile.
SHARED_DESIGN = {"rule-45-highlighted-text", "rule-46-tracked-changes", "rule-55-ellipsis"}


def check_package(path: Path) -> dict:
    problems = []
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                problems.append("document part or content types missing")
            bad = archive.testzip()
            if bad:
                problems.append(f"corrupt member {bad}")
            for name in names:
                if name.endswith(".xml") and b'="NaN"' in archive.read(name):
                    problems.append(f"NaN attribute value in {name}")
        Document(str(path))
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        problems.append(f"cannot open: {error}")
    return {"pass": not problems, "problems": problems}


def source_headings(text: str) -> list[str]:
    found, fenced = [], False
    for line in text.splitlines():
        if line.lstrip().startswith(("```", "~~~")):
            fenced = not fenced
            continue
        match = None if fenced else HEADING.match(line)
        if match:
            found.append(re.sub(r"[*_`]", "", match.group(2)).strip())
    return found


def check_content(path: Path, markdown: str | None) -> dict:
    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    blanks = sorted({m.group(0) for text in paragraphs + cells for m in BLANK.finditer(text)})
    leaked = [p for p in paragraphs[:8] if LEAK.match(p)]
    # A heading may stand inside a side-by-side columns table, so cells count as text too.
    seen = docx_layout.norm("".join(paragraphs + cells))
    missing = [h for h in source_headings(markdown)] if markdown else []
    missing = [h for h in missing if docx_layout.norm(h) not in seen]
    chars = sum(map(len, paragraphs)) + sum(map(len, cells))
    return {"pass": chars > 0 and not blanks and not leaked and not missing, "characters": chars,
            "blanks": blanks[:20], "frontmatter_leak": leaked[:4], "missing_headings": missing[:20]}


def lint(source: Path | None, path: Path, publisher: str | None, kind: str) -> dict:
    registry = SKILL / "templates" / "registry.yaml"
    profile = slop_lint.load_publisher(registry, publisher or "korean-generic")
    rules = slop_lint.load_phrase_rules(SKILL / "references" / "slop_phrase_list.yaml")
    prose, locale = [], "auto"
    if source is not None:
        genre = "manuscript" if kind == "manuscript" else "general"
        prose, _front, _body, locale = slop_lint.lint_text(source.read_text(encoding="utf-8"), profile, rules, "auto", genre)
    design = slop_lint.audit_docx_design(path, profile)
    if publisher is None:
        design = [f for f in design if f.rule_id in SHARED_DESIGN]
    craft = slop_lint.audit_docx_craft(path)
    if locale in {"ko", "mixed"} or source is None:
        design += slop_lint.audit_docx_cjk(path)
    row = lambda f: {"rule": f.rule_id, "severity": f.severity, "line": f.line, "message": f.message, "excerpt": f.excerpt[:80]}  # noqa: E731
    blocking = [row(f) for f in prose + design + craft if f.severity == "HIGH"]
    return {"pass": not blocking, "locale": locale, "blocking": blocking[:40],
            "advisories": [row(f) for f in prose + design + craft if f.severity != "HIGH"][:40]}


def render(path: Path, out: Path) -> Path | None:
    """PDF through LibreOffice with a private profile, so a running office session is never touched."""
    soffice = shutil.which("soffice")
    if not soffice:
        return None
    profile = out / "profile"
    subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                    "--outdir", str(out), str(path)], capture_output=True, text=True, timeout=180)
    pdf = out / f"{path.stem}.pdf"
    return pdf if pdf.exists() else None


def output_checks(path: Path, source: Path | None, args) -> dict:
    front = parse_frontmatter(source.read_text(encoding="utf-8"))[0] if source else {}
    tonality = bool(args.tonality or front.get("tonality"))
    info = docx_layout.read_docx(path)
    found = docx_layout.structural(info, front, args.kind == "manuscript", tonality) + docx_layout.restraint(path, tonality)
    result = {"components": info["components"], "rendered": False}
    if args.compare:
        mine, other = docx_layout.structure(path), docx_layout.structure(args.compare.expanduser().resolve())
        differ = docx_layout.structure_diff(mine, other)
        result["structure"] = {"differ": differ}
        if len(differ) < 3:
            found.append(docx_layout.finding("tonality.structure", "FAIL", None,
                                             f"{len(differ)} structural difference(s): {', '.join(differ) or 'none'}",
                                             "Tonalities differ in structure: title block, summary, numbering, components, running head, contents."))
    if args.layout:
        with tempfile.TemporaryDirectory(prefix="litgrok-docx-gate-") as tmp:
            try:
                pdf = render(path, Path(tmp))
            except subprocess.TimeoutExpired:
                pdf = None
            if pdf is None:
                result["render"] = "LibreOffice is unavailable or the render failed; the page checks did not run"
            else:
                paged, pages = docx_layout.paged(info, docx_layout.read_pdf(pdf), tonality, front)
                found += paged
                result.update(pages, rendered=True)
    result["findings"] = [f for f in found if f["severity"] == "FAIL"]
    result["advisories"] = [f for f in found if f["severity"] != "FAIL"]
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Pass/fail gate for a converted DOCX.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--source", type=Path, help="the Markdown the DOCX was converted from")
    parser.add_argument("--publisher", help="the publisher profile used (korean-generic, elsevier, acs, ieee, nature); omit otherwise")
    parser.add_argument("--kind", choices=("report", "manuscript"), default="report")
    parser.add_argument("--tonality", help="the tonality used, when the source frontmatter does not name it")
    parser.add_argument("--layout", action="store_true", help="render with LibreOffice and run the page checks")
    parser.add_argument("--compare", type=Path, help="the same source in another tonality")
    args = parser.parse_args(argv)

    path = args.docx.expanduser().resolve()
    source = args.source.expanduser().resolve() if args.source else None
    package = check_package(path)
    empty = {"pass": False}
    content = check_content(path, source.read_text(encoding="utf-8") if source else None) if package["pass"] else empty
    lint_report = lint(source, path, args.publisher, args.kind) if package["pass"] else empty
    output = output_checks(path, source, args) if package["pass"] else {"findings": [], "advisories": []}

    reasons = []
    if not package["pass"]:
        reasons.append("package: " + "; ".join(package["problems"][:5]))
    if content.get("blanks"):
        reasons.append(f"content: unfilled blanks {content['blanks'][:5]} — write realistic values and label them as examples")
    if content.get("frontmatter_leak"):
        reasons.append("content: frontmatter printed as body text")
    if content.get("missing_headings"):
        reasons.append(f"content: headings missing from the document {content['missing_headings'][:5]}")
    if package["pass"] and not content.get("characters"):
        reasons.append("content: the document has no text")
    if lint_report.get("blocking"):
        reasons.append("lint: " + ", ".join(sorted({f["rule"] for f in lint_report["blocking"]})))
    if output["findings"]:
        reasons.append("output: " + ", ".join(sorted({f["check"] for f in output["findings"]})))
    report = {"file": str(path), "pass": not reasons, "failure_reasons": reasons, "profile": args.publisher or "plain",
              "kind": args.kind, "package": package, "content": content, "lint": lint_report, "output": output}
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    sys.exit(main())
