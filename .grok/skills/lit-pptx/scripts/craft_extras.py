"""Deterministic slide craft checks, measured from python-pptx and OOXML."""
from __future__ import annotations

import colorsys
import math
import re

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

EMU = 914400
LIMITS = {
    "full_bleed": .92, "accent_saturation": .12, "accent_light_min": .12,
    "accent_light_max": .93, "hue_band": 20, "latin_chars": 90,
    "cjk_chars": 38, "narrow_chars": 10, "group_ratio": 2,
    "symmetric_inset": .06, "radius_tolerance": .02,
    "empty_high": .35, "empty_medium": .25, "rule": .03,
}


# Advance widths of the bundled Pretendard faces in em (the wider of Regular and Bold per class),
# so a measure or a wrap estimate follows the face the deck is set in.
EM = {"hangul": .87, "upper": .67, "lower": .54, "digit": .62, "space": .25, "punct": .41}
# Slides that hold a statement rather than a body: their space is part of the design.
DISPLAY_FAMILIES = re.compile(r"family@(?:cover|section|statement|quote|closing-statement)")
DISPLAY_TITLES = {"title@statement", "title@cover", "title@section"}


def text_em(text):
    """Width of a run of text in em of the bundled face."""
    total = 0.0
    for c in text:
        if "\uac00" <= c <= "\ud7af" or "\u3130" <= c <= "\u318f" or "\u4e00" <= c <= "\u9fff":
            total += EM["hangul"]
        elif c.isspace():
            total += EM["space"]
        elif c.isdigit():
            total += EM["digit"]
        elif c.isupper():
            total += EM["upper"]
        elif c.isalpha():
            total += EM["lower"]
        else:
            total += EM["punct"]
    return total


def named_title(slide):
    """The frame the deck engine names as the slide title (`title@<treatment>`), if any."""
    return next((s for s in slide.shapes if s.name.startswith("title@") and getattr(s, "has_text_frame", False)), None)


def display_slide(slide):
    """A cover, section, statement or quote slide drawn by the deck engine."""
    return any(DISPLAY_FAMILIES.match(s.name) or s.name in DISPLAY_TITLES for s in slide.shapes)


def _box(shape):
    return tuple(float(getattr(shape, key)) / EMU for key in ("left", "top", "width", "height"))


def _distance(a, b):
    return math.hypot(max(0, a[0] - b[0] - b[2], b[0] - a[0] - a[2]),
                      max(0, a[1] - b[1] - b[3], b[1] - a[1] - a[3]))


def _hue(shape):
    try:
        rgb = shape.fill.fore_color.rgb
        if rgb is None:
            return None
        vals = tuple(int(str(rgb)[i:i + 2], 16) / 255 for i in (0, 2, 4))
        hue, sat, light = colorsys.rgb_to_hls(*vals)
        if sat < LIMITS["accent_saturation"] or not LIMITS["accent_light_min"] <= light <= LIMITS["accent_light_max"]:
            return None
        return hue * 360
    except (AttributeError, TypeError, ValueError):
        return None


def _display(slide):
    text = [s for s in slide.shapes if getattr(s, "has_text_frame", False) and s.text.strip()]
    if len(text) > 2 or sum(len(s.text.strip()) for s in text) > 80 or any(
        getattr(s, "has_table", False) or s.shape_type in (MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.AUTO_SHAPE)
        for s in slide.shapes
    ):
        return False
    sizes = [r.font.size.pt for s in text for p in s.text_frame.paragraphs for r in p.runs if r.font.size]
    return bool(sizes and max(sizes) >= 30)


def _round_radius(shape):
    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or shape.auto_shape_type != MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
        return None
    match = re.search(r'<a:gd[^>]*name="adj"[^>]*fmla="val\s+(\d+)"', shape._element.xml)
    adj = int(match.group(1)) / 100000 if match else .16667
    return adj * min(_box(shape)[2:])


def assess(prs):
    findings = []

    def add(rule, severity, slide, value, threshold, tier="measured", note=""):
        findings.append({"rule": rule, "severity": severity, "slide": slide,
                         "tier": tier, "value": value, "threshold": threshold, "note": note})

    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    for number, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        display = _display(slide)
        hue_families = []
        for shape in shapes:
            box = _box(shape)
            if box[2] * box[3] >= LIMITS["full_bleed"] * W * H:
                continue
            if shape.name.startswith("lit-notice") or min(box[2], box[3]) < LIMITS["rule"]:
                continue
            hue = _hue(shape)
            if hue is not None and not any(min(abs(hue - old), 360 - abs(hue - old)) <= LIMITS["hue_band"] for old in hue_families):
                hue_families.append(hue)
        count = len(hue_families)
        high = 4 if display else 3
        medium = 3 if display else 2
        if count >= medium:
            add("OF-101", "HIGH" if count >= high else "MEDIUM", number, count, {"medium": medium, "high": high})

        text_shapes = [s for s in shapes if getattr(s, "has_text_frame", False) and s.text.strip()]
        title = named_title(slide) or (min(text_shapes, key=lambda s: _box(s)[1]) if text_shapes else None)
        # A cover or section page the engine names is a display page: its preview lines keep their own measure
        # (two lines an opened slide, within 70 % of the page), so the body measure does not apply there.
        family = next((s.name.split("@", 1)[1] for s in shapes if s.name.startswith("family@")), "")
        for shape in text_shapes:
            if shape is title or re.match(r"(cover|section)(-|$)", family):
                continue
            sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
            pt = max(sizes) if sizes else 18
            # Body text only: captions and table notes at 11 pt run as wide as the visual they label.
            if not 11 < pt < 30:
                continue
            content = shape.text.strip()
            frame = max(1, shape.width / EMU * 72)
            # Each paragraph or explicit break is a line of its own that may wrap again in the frame.
            segments = [seg.strip() for seg in re.split(r"[\n\v]", content) if seg.strip()]
            wrapped = [(seg, max(1, math.ceil(text_em(seg) * pt / frame))) for seg in segments]
            lines = sum(n for _, n in wrapped)
            if lines < 2:
                continue
            cjk = sum("\uac00" <= c <= "\ud7af" for c in content) >= .5 * len(content.replace(" ", ""))
            # A Korean measure counts glyphs, so the spaces between 어절 are not counted.
            measure = max(len(seg.replace(" ", "") if cjk else seg) / n for seg, n in wrapped)
            ceiling = LIMITS["cjk_chars"] if cjk else LIMITS["latin_chars"]
            if measure > ceiling:
                add("OF-102", "HIGH", number, round(measure, 1), ceiling, "derived", shape.name)
            elif lines >= 3 and measure < LIMITS["narrow_chars"]:
                add("OF-102", "MEDIUM", number, round(measure, 1), LIMITS["narrow_chars"], "derived", shape.name)

        for shape in shapes:
            if not getattr(shape, "has_table", False):
                continue
            table = shape.table
            if len(table.rows) < 3:
                continue
            for col in range(len(table.columns)):
                cells = [table.cell(row, col) for row in range(1, len(table.rows))]
                numeric = [c for c in cells if re.fullmatch(r"[\d.,%+\-±()$€£¥▲▼\s]+", c.text.strip()) and re.search(r"\d", c.text)]
                if len(numeric) / len(cells) < .7:
                    continue
                for cell in numeric:
                    for paragraph in cell.text_frame.paragraphs:
                        if paragraph.alignment is None:
                            add("OF-103", "MEDIUM", number, cell.text, "right alignment", "derived", "inherited table alignment")
                        elif paragraph.alignment != PP_ALIGN.RIGHT:
                            add("OF-103", "HIGH", number, cell.text, "right alignment")

        rounded = [(s, _box(s), _round_radius(s)) for s in shapes]
        rounded = [(s, b, r) for s, b, r in rounded if r is not None]
        for outer, ob, outer_r in rounded:
            for inner, ib, inner_r in rounded:
                if inner is outer:
                    continue
                insets = [ib[0] - ob[0], ib[1] - ob[1], ob[0] + ob[2] - ib[0] - ib[2], ob[1] + ob[3] - ib[1] - ib[3]]
                if min(insets) < 0 or max(insets) - min(insets) > LIMITS["symmetric_inset"]:
                    continue
                expected = max(0, outer_r - sum(insets) / 4)
                if outer_r > .04 and abs(inner_r - expected) > LIMITS["radius_tolerance"]:
                    add("OF-104", "HIGH", number, round(inner_r, 3), round(expected, 3), note=f"{outer.name}/{inner.name}")

        cards = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and .04 * W * H < _box(s)[2] * _box(s)[3] < LIMITS["full_bleed"] * W * H and _hue(s) is not None]
        groups = []
        for card in cards:
            x, y, w, h = _box(card)
            children = [s for s in shapes if s is not card and x <= _box(s)[0] and y <= _box(s)[1] and _box(s)[0] + _box(s)[2] <= x + w and _box(s)[1] + _box(s)[3] <= y + h]
            if len(children) >= 2:
                groups.append((card, children))
        if len(groups) >= 2:
            intra = [_distance(_box(a), _box(b)) for _, children in groups for i, a in enumerate(children) for b in children[i + 1:]]
            inter = [_distance(_box(a), _box(b)) for i, (a, _) in enumerate(groups) for b, _ in groups[i + 1:]]
            if intra and inter and min(inter) < LIMITS["group_ratio"] * max(intra):
                add("OF-105", "MEDIUM", number, round(min(inter), 3), round(LIMITS["group_ratio"] * max(intra), 3))

        for shape in shapes:
            xml = shape._element.xml
            if shape._element.xpath('.//a:rPr/a:gradFill'):
                add("OF-106", "HIGH", number, shape.name, "no gradient text")
            if "<a:glow" in xml:
                add("OF-107", "HIGH", number, shape.name, "no glow")
            for char in re.findall(r'<a:buChar[^>]*char="([^"]+)"', xml):
                if any(0x1F000 <= ord(c) <= 0x1FAFF or 0x2600 <= ord(c) <= 0x27BF for c in char):
                    add("OF-108", "HIGH", number, char, "no emoji bullet")

        if not display and not display_slide(slide) and number > 1 and text_shapes:
            title_bottom = _box(title)[1] + _box(title)[3] if title else H * .18
            top, bottom = title_bottom + .1, H - .55

            def content(s):
                return (getattr(s, "has_table", False) or getattr(s, "has_chart", False) or s.shape_type == MSO_SHAPE_TYPE.PICTURE
                        or getattr(s, "has_text_frame", False) and s.text.strip())

            # A title on the body floor closes the page, so the body is the zone above it. A side title's body
            # is the zone beside its rail from the title's top; what stands in the rail under it is the rail's.
            rail = set()
            named = title is not None and title.name.startswith("title@")
            if named and _box(title)[1] > H * .5:
                top, bottom = .5, _box(title)[1] - .1
            elif named and _box(title)[0] <= W * .1 and _box(title)[2] <= W * .35:
                edge = _box(title)[0] + _box(title)[2]
                others = [s for s in shapes if s is not title and not s.name.startswith("lit-notice") and content(s) and _box(s)[1] < bottom]
                under = [s for s in others if _box(s)[0] + _box(s)[2] <= edge + .05 and _box(s)[1] >= title_bottom - .05]
                if any(_box(s)[0] >= edge - .05 for s in others) and all(s in under or _box(s)[0] >= edge - .05 for s in others):
                    top, rail = _box(title)[1], {id(s) for s in under}

            # The empty band is a height at which nothing on the slide holds content: blocks are clipped to
            # the body and joined, so a short label beside a tall chart, a column that starts beside a side
            # title, or a picture behind the body all count for the height they cover.
            spans = sorted((max(top, _box(s)[1]), min(bottom, _box(s)[1] + _box(s)[3])) for s in shapes
                           if s is not title and id(s) not in rail and not s.name.startswith("lit-notice") and content(s)
                           and _box(s)[1] < bottom and _box(s)[1] + _box(s)[3] > top)
            if spans and bottom > top:
                cursor = top
                gaps = []
                for start, end in spans:
                    gaps.append(max(0, start - cursor))
                    cursor = max(cursor, end)
                gaps.append(max(0, bottom - cursor))
                fraction = max(gaps) / (bottom - top)
                if fraction >= LIMITS["empty_medium"]:
                    add("OF-109", "HIGH" if fraction > LIMITS["empty_high"] else "MEDIUM", number, round(fraction, 3), LIMITS["empty_high"])

    return {"pass": not any(f["severity"] == "HIGH" for f in findings), "findings": findings[:100]}
