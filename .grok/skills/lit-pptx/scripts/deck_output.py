#!/usr/bin/env python3
"""Deck-wide output checks, read from the compiled file.

  OF-110  a deck of 8+ slides shows at least three title treatments, each drawn as it is named
  OF-111  content slides do not repeat one composition (title zone x body partition x dominant content)
  OF-112  the median empty band under the body stays at or under 0.20
  OF-113  every slide under one title treatment keeps that treatment's frame
  OF-114  titles and the cover subtitle are topic labels, never sentences
  OF-115  no region stands empty beside a full one: the floor under a bottom title, the rail under a
          side title, the column beside a visual, a title panel over a picture, the page above a plate
  OF-116  (qa_deck.py --sibling) another tonality of the same source draws another skeleton
  OF-117  no mostly-white figure picture on a dark slide ground
  OF-118  a bold run-in label is followed by its separator ("요약: …", "Summary: …")
  OF-119  an agenda's title carries no count numeral beside it (its rows carry the numbers)

The deck engine names its title frame `title@<treatment>` and the first shape of a slide
`family@<layout family>`. A name is only compared with what the geometry shows; it is never the
evidence on its own. A deck without those names was not drawn by the engine and is out of scope.
A legacy template (`title@legacy`) asked for one look on purpose, so its variety and band findings
are MEDIUM. Geometry is in points on the 960 x 540 or 720 x 540 canvas.
"""
from __future__ import annotations

import io
import json
import math
import re
import statistics
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from craft_extras import text_em

PT = 12700
TOL, SIZE_TOL = 2.0, 0.5
FLOOR, FOOTER = 486.0, 492.0
LIMIT = {"slides": 8, "treatments": 3, "share": 0.40, "cap": 5, "band": 0.20, "anchor": 1.0}
# Per-slide band caps by density: (highest density in the step, cap).
BAND_CAPS = ((2, .28), (4, .24), (6, .20), (8, .16), (10, .12))
ZONES = {"top-rule": "top", "top-plain-large": "top", "kicker-numeral": "top", "side-rail": "side",
         "band": "band", "overlay": "overlay", "bottom-anchor": "bottom", "statement": "none"}
# The widest four-column title and the earliest column-5 start over all density grids.
RAIL = {960: (288.0, 336.0), 720: (216.0, 252.0)}
MARGINS = {960: {"airy": 60, "standard": 48, "dense": 36, "compact": 24}, 720: {"airy": 54, "standard": 42, "dense": 30, "compact": 24}}
STAMP = re.compile(r"tonality=(\S+) density=(\d+) grid=(\w+) variance=(\d+)")
NUMERAL = re.compile(r"\d+(?:[.:/-]\d+)?\.?")
BAND_EXEMPT = re.compile(r"(?:cover|section|statement|quote|closing-statement)")


def _shape(shape):
    if shape.left is None or shape.width is None:
        return None
    text = shape.text_frame.text.strip() if getattr(shape, "has_text_frame", False) and shape.has_text_frame else ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        kind = "picture"
    elif getattr(shape, "has_table", False) and shape.has_table:
        kind = "table"
    elif getattr(shape, "has_chart", False) and shape.has_chart:
        kind = "chart"
    else:
        kind = "text" if text else "shape"
    x, y, w, h = (v / PT for v in (shape.left, shape.top, shape.width, shape.height))
    if kind == "table":
        h = max(h, sum(row.height for row in shape.table.rows) / PT)
    sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs
             if r.text.strip() and r.font.size is not None] if text else []
    try:
        filled = kind == "shape" and shape.fill.type == 1
    except (AttributeError, TypeError):
        filled = False
    return {"name": shape.name, "kind": kind, "x": x, "y": y, "w": w, "h": h, "text": text, "sizes": sizes, "filled": filled}


def _inside(e, box, tol=1.0):
    cx, cy = e["x"] + e["w"] / 2, e["y"] + e["h"] / 2
    return box["x"] - tol <= cx <= box["x"] + box["w"] + tol and box["y"] - tol <= cy <= box["y"] + box["h"] + tol


def elements(slide, W, H):
    """Drawn shapes in points, each with a role: content, card (a fill that holds text) or decoration."""
    out = [e for e in map(_shape, slide.shapes) if e and e["w"] + e["h"] > 0 and not e["name"].startswith("lit-notice")]
    for e in out:
        if e["kind"] != "shape":
            e["role"] = "content"
            continue
        bleeds = e["x"] <= 1 or e["y"] <= 1 or e["x"] + e["w"] >= W - 1 or e["y"] + e["h"] >= H - 1
        holds = any(o["kind"] != "shape" and _inside(o, e) for o in out if o is not e)
        card = e["filled"] and holds and not bleeds and e["w"] >= 100 and e["h"] >= 40 and e["w"] * e["h"] < .5 * W * H
        e["role"] = "card" if card else "decoration"
    return out


def treatment(title, els, W, H):
    """The title treatment the geometry shows; the tests run in a fixed order and the first match wins."""
    size = max(title["sizes"] or [0])
    bottom = title["y"] + title["h"]
    others = [e for e in els if e is not title]
    body = [e for e in others if e["y"] < FOOTER and e["role"] != "decoration"]

    def under_title(e):
        return (e["x"] - TOL <= title["x"] and title["x"] + title["w"] <= e["x"] + e["w"] + TOL
                and e["y"] - TOL <= title["y"] and bottom <= e["y"] + e["h"] + TOL)

    if (any(e["kind"] == "picture" and e["w"] * e["h"] >= .8 * W * H for e in others) and title["y"] >= 300 - TOL
            and any(e["kind"] == "shape" and e["filled"] and under_title(e) for e in others)):
        return "overlay"
    if any(e["kind"] == "shape" and e["filled"] and e["y"] <= 2 + TOL and 96 - TOL <= e["h"] <= 120 + TOL
           and e["w"] >= .9 * W and _inside(title, e, TOL) for e in others):
        return "band"
    if (size >= 36 - SIZE_TOL and 220 - TOL <= title["y"] + title["h"] / 2 <= 380 + TOL
            and not any(e["kind"] in ("table", "chart", "picture") or e["role"] == "card" for e in body)
            and sum(e["kind"] == "text" for e in body) <= 1):
        return "statement"
    if title["y"] >= 360 - TOL and all(e["y"] <= title["y"] + TOL for e in body if e["kind"] == "text"):
        return "bottom-anchor"
    rail, column5 = RAIL.get(round(W), RAIL[960])
    beside = [e for e in body if e["x"] + e["w"] / 2 > title["x"] + title["w"]]
    # What stands in the rail under the title (criteria labels, takeaways, the source) is the rail's; only a
    # block that runs from the rail on into the body makes the title something else.
    if (title["w"] <= rail + TOL and beside and min(e["x"] for e in beside) >= column5 - TOL
            and not any(e["x"] < column5 - TOL and e["y"] > bottom and e["x"] + e["w"] > column5 + TOL for e in body)):
        return "side-rail"
    if any(e["kind"] == "text" and e["sizes"] and NUMERAL.fullmatch(e["text"]) and max(e["sizes"]) >= 44 - SIZE_TOL
           and e["x"] + e["w"] <= title["x"] + TOL and e["y"] >= 30 - TOL and e["y"] + e["h"] <= 114 + TOL for e in others):
        return "kicker-numeral"
    if title["y"] <= 48 + TOL:
        # A rule is a line under the title, never a list-marker dot.
        rule = any(e["kind"] == "shape" and e["h"] <= 8 and e["w"] >= max(16, 4 * e["h"]) and bottom - TOL <= e["y"] <= bottom + 24
                   and e["x"] < title["x"] + title["w"] and e["x"] + e["w"] > title["x"] for e in others)
        if size >= 36 - SIZE_TOL and not rule:
            return "top-plain-large"
        if size <= 32 + SIZE_TOL and rule:
            return "top-rule"
    return "unclassified"


def _figure(e):
    """A run of 44 pt or more that is mostly digits reads as a number, not as text."""
    compact = re.sub(r"\s+", "", e["text"])
    return (e["kind"] == "text" and bool(e["sizes"]) and max(e["sizes"]) >= 44 - SIZE_TOL
            and len(re.findall(r"\d", compact)) >= .5 * max(1, len(compact)))


def signature(title, kind, els, W, H, margin):
    """(title zone, body partition, dominant content) of one slide."""
    zone = ZONES.get(kind, "none")
    if kind == "statement":
        return (zone, "one", "text")
    column5 = RAIL.get(round(W), RAIL[960])[1]
    body = []
    for e in els:
        if e is title or e["y"] >= FOOTER or e["role"] == "decoration":
            continue
        if kind == "side-rail" and e["x"] + e["w"] <= column5 + TOL:
            continue  # the rail under a side title is part of the title zone, not a body column
        if e["kind"] == "shape" and title is not None and _inside(title, e, 0):
            continue  # the band, rail or panel that carries the title
        if kind == "kicker-numeral" and e["kind"] == "text" and NUMERAL.fullmatch(e["text"]) and e["y"] + e["h"] <= 116:
            continue
        body.append(e)
    cards = [e for e in body if e["role"] == "card"]
    blocks = [e for e in body if not any(c is not e and _inside(e, c) for c in cards)]
    if any(e["kind"] == "picture" and e["w"] * e["h"] >= .8 * W * H for e in els):
        partition = "full-bleed"
    elif not blocks:
        partition = "one"
    else:
        # Columns join boxes whose x-extents overlap by half the narrower one; a box wider than 70 % of
        # the body (a takeaway under two columns) never joins columns.
        narrow = [e for e in blocks if e["w"] < .7 * (W - 2 * margin)] or blocks
        columns = []
        for e in sorted(narrow, key=lambda e: e["x"]):
            hits = [c for c in columns if min(c[1], e["x"] + e["w"]) - max(c[0], e["x"]) >= .5 * min(c[1] - c[0], e["w"])]
            merged = [min([e["x"]] + [c[0] for c in hits]), max([e["x"] + e["w"]] + [c[1] for c in hits]),
                      [e] + [i for c in hits for i in c[2]]]
            columns = [c for c in columns if c not in hits] + [merged]
        column = {id(i): n for n, c in enumerate(columns) for i in c[2]}
        rows, grid_rows = [], 0
        for e in sorted(narrow, key=lambda e: e["y"]):
            if rows and e["y"] < rows[-1][0] - 2:
                rows[-1][0] = max(rows[-1][0], e["y"] + e["h"])
                rows[-1][1].append(e)
            else:
                rows.append([e["y"] + e["h"], [e]])
        grid_rows = sum(len({column[id(i)] for i in r[1]}) >= 2 for r in rows)
        k = len(columns)
        if k >= 4 or (k >= 2 and grid_rows >= 2):
            partition = "grid"
        elif k == 3:
            partition = "three"
        elif k == 2:
            a, b = (c[1] - c[0] for c in columns)
            partition = "two-even" if min(a, b) / max(a, b) >= .85 else "two-asym"
        else:
            partition = "one"
    area = {}
    for e in blocks:
        key = {"picture": "image", "table": "table", "chart": "chart"}.get(e["kind"])
        if key is None and e["kind"] == "text":
            key = "number" if _figure(e) else "text"
        if key:
            area[key] = area.get(key, 0) + e["w"] * e["h"]
    for c in cards:
        key = "number" if any(_figure(e) for e in body if e is not c and _inside(e, c)) else "text"
        area[key] = area.get(key, 0) + c["w"] * c["h"]
    return (zone, partition, max(area, key=area.get) if area else "text")


def band(title, els):
    """Share of the body left empty under its lowest block, from the title (or the top margin when
    the title sits low) down to the body floor. A title set at the foot is ink there too."""
    top = title["y"] + title["h"] if title and title["y"] + title["h"] < FLOOR * .6 else 36.0
    if FLOOR - top < 10:
        return None
    lowest = max((min(FLOOR, e["y"] + e["h"]) for e in els if e["role"] != "decoration" and e["y"] < FOOTER
                  and e["y"] < FLOOR and e["y"] + e["h"] > top), default=top)
    return max(0.0, (FLOOR - lowest) / (FLOOR - top))


def _role(family):
    return next((r for r in ("cover", "section", "closing") if family == r or family.startswith(r + "-")), "content")


def _most_treatments(variance):
    return 2 if variance <= 3 else 3 if variance == 4 else 4 if variance <= 6 else 5


# Titles are labels. Korean is read on its last 어절: a sentence ending or a nominalised claim makes
# the title a sentence. English is read for closing punctuation, an auxiliary or a finite report
# verb. A closing parenthesis is set aside first; a question or a quotation is not a claim.
KO_ENDING = re.compile(r"(?:니다|[아어여해게네지세래]요|죠|[았었였했됐겠]다|[가-힣]다|[가-힣][함됨]|[있없]음|[았었였했됐]음)$")
KO_NOUNS = {"바다", "판다", "소다", "람다", "캐나다", "어젠다", "아젠다", "보다", "다", "포함", "함"}
EN_AUX = re.compile(r"\b(?:is|are|was|were|has|have|had|will|would|can|could|should|must|does|did|isn't|aren't|wasn't|won't|doesn't|didn't)\b", re.I)
# Finite verbs that rarely stand as nouns in a label: past forms, and a few third-person forms whose
# plural-noun reading is rare. Forms such as "cuts", "leads", "works" or "won" read as nouns too often.
EN_FINITE = {
    "rose", "grew", "fell", "climbed", "dropped", "doubled", "tripled", "halved", "increased", "decreased",
    "declined", "improved", "reached", "exceeded", "missed", "outperformed", "drove", "showed", "proved",
    "stayed", "remained", "needed", "made", "took", "saved", "lifted", "slowed", "failed", "worked", "explained",
    "outperforms", "exceeds", "beats", "proves", "explains", "improves", "grows", "slows", "fails", "loses", "pays",
}


def sentence(text):
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    t = re.sub(r"\s*\([^()]*\)$", "", t).strip().rstrip("\"'”’")
    if not t or t.endswith("?") or t[0] in "“\"「『‘'":
        return False
    if re.search(r"[가-힣]", t):
        last = re.sub(r"[.!…]+$", "", t.split(" ")[-1])
        return last not in KO_NOUNS and bool(KO_ENDING.search(last))
    if t[-1] in ".!":
        return True
    words = re.findall(r"[A-Za-z][A-Za-z'-]*", t)  # a hyphenated compound ("cut-over") is one word
    return bool(EN_AUX.search(t)) or any(w.lower() in EN_FINITE for w in words[1:])


# OF-115. The band (OF-112) reads the body as one page, so a region can still stand empty beside a full one.
# Each reading is a share of its own region, or under a bottom title a distance from the floor.
REGION = {"floor_gap": 12.0, "share": 0.40}


def _finding(rule, slide, value, threshold, note):
    return {"rule": rule, "severity": "HIGH", "slide": slide, "tier": "measured", "value": value, "threshold": threshold, "note": note}


def _title_lines(title):
    """Lines a title is set in: its explicit breaks, each wrapped again in the frame at the bundled face's widths."""
    size = max(title["sizes"] or [0]) or 1
    segments = [seg for seg in re.split(r"[\n\v]", title["text"]) if seg.strip()]
    return sum(max(1, math.ceil(text_em(seg) * size / max(1.0, title["w"]))) for seg in segments) or 1


def _largest_gap(top, bottom, spans):
    """The tallest stretch of [top, bottom] that no (y0, y1) span covers."""
    gap, cursor = 0.0, top
    for y0, y1 in sorted(spans):
        if y1 <= cursor:
            continue
        gap = max(gap, max(0.0, min(y0, bottom) - cursor))
        cursor = max(cursor, y1)
    return max(gap, bottom - cursor)


def _short_column(content, title, kind, body_floor, column5, fields, family=None):
    """A column of a two-column body that stops short of its neighbour: the largest empty band inside it, from
    the body's top down to where the neighbour ends, is over 40 % of the body. A source or note strip that closes
    the column does not hide the band above it. Beside the figures of a kpi-row or big-number slide the column is
    measured to the body floor (spec-v2 Amendments 9)."""
    # The body starts under a top title; beside a side title (in its rail) or over a bottom title it starts at the top.
    top_of_title = title["y"] + title["h"] if kind not in ("bottom-anchor", "side-rail") else 0.0
    body = [e for e in content if e["y"] >= top_of_title - TOL and e["y"] < body_floor and e["w"] > 0 and e["h"] > 0
            and not (kind == "side-rail" and e["x"] + e["w"] <= column5 + TOL) and not any(_inside(e, f) for f in fields)]
    if len(body) < 2:
        return None
    left_edge, right_edge = min(e["x"] for e in body), max(e["x"] + e["w"] for e in body)
    for cut in sorted({e["x"] + e["w"] for e in body}):
        a = [e for e in body if e["x"] + e["w"] <= cut + TOL]
        b = [e for e in body if e["x"] >= cut - TOL]
        # A block across most of the body (a note under both columns) belongs to neither column; a narrower
        # block across the cut means there is no column line here.
        across = [e for e in body if e not in a and e not in b]
        if not a or not b or any(e["w"] < .7 * (right_edge - left_edge) for e in across):
            continue
        top = min(e["y"] for e in a + b)
        for column, other in ((a, b), (b, a)):
            # Takeaways beside a chart that close with the chart's values: the table counts at its drawn extent.
            values = any(e["kind"] == "chart" for e in other) and any(e["kind"] == "text" for e in column) \
                and not any(e["kind"] in ("chart", "picture") for e in column)
            # A text column, or a picture drawn smaller than its column, beside a column that runs on.
            if not any(e["kind"] in ("text", "picture") for e in column) or (any(e["kind"] in ("chart", "table") for e in column) and not values):
                continue
            # Row heads (a matrix's or a comparison's short labels, each level with its row) are no column to fill.
            if all(e["kind"] == "text" and e["h"] <= 56 and len(re.sub(r"\s+", "", e["text"])) <= 16 for e in column):
                continue
            # Read whatever the neighbour's length: one that stops short still marks how far this column should run.
            reach = min(body_floor, max(e["y"] + e["h"] for e in other))
            # A takeaway column beside the figures of a kpi-row or big-number slide runs to the body floor: the rows stop
            # where their spacing ends. The figures' own labels stand exactly level with figures below the first row.
            size = lambda side: max([max(e["sizes"]) for e in side if e["kind"] == "text" and e["sizes"]] or [0.0])  # noqa: E731
            figures = [e for e in other if e["kind"] == "text" and e["sizes"] and max(e["sizes"]) > size(column) + 4]
            labels = any(abs(e["y"] - f["y"]) <= .5 and f["y"] > top + TOL for e in column for f in figures)
            if family in ("kpi-row", "big-number") and figures and not labels:
                reach = body_floor
            share = _largest_gap(top, reach, [(e["y"], e["y"] + e["h"]) for e in column]) / max(1.0, body_floor - top)
            if share > REGION["share"]:
                what = "picture" if any(e["kind"] == "picture" for e in column) else "text"
                return _finding("OF-115", None, round(share, 3), REGION["share"],
                                f"the {what} column beside a longer one leaves its share of the body empty inside it: fill it by layout "
                                "(lead size, the source and values in the column, a narrower column with a wider visual, or the points under the visual)")
    return None


def empty_regions(slides, W, H):
    found = []
    column5 = RAIL.get(round(W), RAIL[960])[1]
    full = lambda e: e["kind"] == "picture" and e["w"] * e["h"] >= .8 * W * H  # noqa: E731
    for s in slides:
        title, kind = s["title"], s["kind"]
        if title is None or not title["sizes"]:
            continue
        content = [e for e in s["els"] if e is not title and e["role"] != "decoration" and e["y"] < FOOTER and not full(e)]
        if s["role"] in ("content", "closing"):
            if kind == "bottom-anchor":
                size = max(title["sizes"])
                gap = FLOOR - (title["y"] + _title_lines(title) * size * 1.15)
                if gap > REGION["floor_gap"]:
                    found.append(_finding("OF-115", s["n"], round(gap, 1), REGION["floor_gap"],
                                          "the bottom title ends above the body floor: stand its last line on the floor"))
            if kind == "side-rail" and s["family"] not in ("quote", "statement"):
                top = title["y"] + title["h"]  # the frame keeps a spare line for a title that wraps once more
                rail = [e for e in content if e["x"] + e["w"] <= column5 + TOL and e["y"] + e["h"] > top]
                share = _largest_gap(top, FLOOR, [(e["y"], e["y"] + e["h"]) for e in rail]) / max(1.0, FLOOR - top)
                if share > REGION["share"]:
                    found.append(_finding("OF-115", s["n"], round(share, 3), REGION["share"],
                                          "the side rail under the title stands empty: give it the criteria, takeaways or source, or take another title"))
            # A takeaway column beside a chart, table or picture that runs on below it. Text on a colour field
            # (a closing's ask) stands on its own ground, not in a column.
            body_floor = title["y"] - TOL if kind == "bottom-anchor" else FLOOR
            fields = [e for e in s["els"] if e["kind"] == "shape" and e["filled"] and e["w"] * e["h"] >= .1 * W * H]
            short = _short_column(content, title, kind, body_floor, column5, fields, s["family"])
            if short:
                found.append({**short, "slide": s["n"]})
            for v in ([] if short else [e for e in content if e["kind"] in ("chart", "table", "picture")]):
                beside = [e for e in content if e["kind"] == "text" and e["y"] < v["y"] + v["h"] and e["y"] + e["h"] > v["y"]
                          and (e["x"] >= v["x"] + v["w"] - TOL or e["x"] + e["w"] <= v["x"] + TOL)
                          and not any(_inside(e, f) for f in fields)]
                if not beside:
                    continue
                left, right = min(e["x"] for e in beside), max(e["x"] + e["w"] for e in beside)
                top = min([v["y"]] + [e["y"] for e in beside])
                column = max(e["y"] + e["h"] for e in content if left - TOL <= e["x"] + e["w"] / 2 <= right + TOL)
                visual = min(body_floor, max(e["y"] + e["h"] for e in content if v["x"] - TOL <= e["x"] + e["w"] / 2 <= v["x"] + v["w"] + TOL))
                share = (visual - column) / max(1.0, body_floor - top)
                if share > REGION["share"]:
                    found.append(_finding("OF-115", s["n"], round(share, 3), REGION["share"],
                                          f"the column beside the {v['kind']} stops short of it: close it with the source and note, the values or the takeaways"))
                    break
        if s["role"] in ("cover", "section") or kind in ("statement", "overlay"):
            # A title panel over a picture, and a plate under a statement or quote. A typographic cover's plate
            # is the cover's own device, not a region its text must fill.
            pictured = any(full(e) for e in s["els"])
            for p in [e for e in s["els"] if e["kind"] == "shape" and e["filled"] and _inside(title, e) and e["w"] * e["h"] < .9 * W * H]:
                inside = [e for e in content + [title] if e is not p and _inside(e, p)]
                if p["y"] > TOL and pictured:
                    floor = min(p["y"] + p["h"], FLOOR)
                    share = _largest_gap(p["y"], floor, [(e["y"], e["y"] + e["h"]) for e in inside]) / max(1.0, floor - p["y"])
                    if share > REGION["share"]:
                        found.append(_finding("OF-115", s["n"], round(share, 3), REGION["share"],
                                              "the panel that carries the title stands partly empty: make it as tall as its text, on the page foot"))
                if p["y"] > TOL and kind == "statement" and p["w"] >= .9 * W and not any(e["y"] + e["h"] <= p["y"] + TOL for e in content):
                    found.append(_finding("OF-115", s["n"], round(p["y"], 1), 0,
                                          "a plate under part of the page leaves the open page above it empty: use the whole-page field"))
    return found


# OF-116. A skeleton is each content slide's title zone and partition, as OF-111 reads them; two tonalities of
# one source whose skeletons match differ in colour only.
def skeleton(prs):
    return [s["sig"][:2] for s in _read(prs) if s["role"] == "content"]


def compare_skeletons(prs, sibling, name="the sibling deck"):
    a, b = skeleton(prs), skeleton(sibling)
    n = min(len(a), len(b))
    differ = sum(x != y for x, y in zip(a, b))
    # Two slides once the deck has four content slides: one differing slide left the alternatives reading as
    # one deck in two colours.
    need = 2 if n >= 4 else 1
    if not n or differ >= need:
        return []
    return [_finding("OF-116", None, differ, need, f"{differ} of {n} content slides differ in title zone or partition from {name} "
                     f"(needs {need}): give the tonality its own structure on more slides")]


# OF-117. A chart or diagram exported on white stands on a dark ground as a white card with small grey labels.
# A figure is a picture of mostly near-white pixels; a photograph rarely is.
FIGURE = {"ground_max": .25, "near_white": 235, "share": .5, "min_area": .03}


def light_figures(prs, W, H):
    found = []
    try:
        from PIL import Image
    except ImportError:
        return found
    for number, slide in enumerate(prs.slides, 1):
        try:
            rgb = slide.background.fill.fore_color.rgb if slide.background.fill.type == 1 else None
        except (AttributeError, TypeError, ValueError):
            rgb = None
        if rgb is None:
            continue
        r, g, b = (int(str(rgb)[i:i + 2], 16) / 255 for i in (0, 2, 4))
        if .2126 * r + .7152 * g + .0722 * b > FIGURE["ground_max"]:
            continue
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or shape.width is None:
                continue
            area = shape.width / PT * shape.height / PT
            if not FIGURE["min_area"] * W * H <= area < .8 * W * H:
                continue
            with Image.open(io.BytesIO(shape.image.blob)) as im:
                raw = im.convert("RGB").resize((64, 64)).tobytes()
            share = sum(min(raw[i:i + 3]) >= FIGURE["near_white"] for i in range(0, len(raw), 3)) / (len(raw) / 3)
            if share > FIGURE["share"]:
                found.append(_finding("OF-117", number, round(share, 3), FIGURE["share"],
                                      f"{shape.name}: a light figure on a dark ground; put its dark variant beside it (name.dark.png) or draw a native chart"))
    return found


# OF-118. A point that opens with a bold label and runs on in regular weight reads as one sentence when nothing
# separates the two ("요약 앱에서 …"): the label takes a colon unless it ends in its own separator.
RUN_IN_SEPARATOR = re.compile(r"[:：.)\]?!–—\-]$")


def run_ins(prs):
    found = []
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame or shape.name.startswith(("title@", "lit-notice")):
                continue
            for paragraph in shape.text_frame.paragraphs:
                runs = [r for r in paragraph.runs if r.text]
                if len(runs) < 2 or not runs[0].font.bold or runs[1].font.bold:
                    continue
                label, rest = runs[0].text, runs[1].text
                if not label.strip() or not re.search(r"\w", rest) or RUN_IN_SEPARATOR.search(label.strip()) or rest.lstrip()[:1] in ":：":
                    continue
                found.append(_finding("OF-118", number, (label + rest)[:40], "a separator",
                                      f"{shape.name}: a bold run-in label without a separator; write the label with its colon (**요약:** …) or let the engine set it"))
    return found


# OF-119. An agenda numbers its rows; a numeral beside its title repeats their count and reads as a stray section number.
def agenda_numerals(slides):
    return [_finding("OF-119", s["n"], s["kind"], "no numeral", "a numeral beside the agenda title repeats the count of its rows: take a title without one")
            for s in slides if s["family"] == "agenda" and s["kind"] == "kicker-numeral"]


def _read(prs):
    """Every slide's elements, family, role, written and drawn treatment and (content slides) signature."""
    W, H = prs.slide_width / PT, prs.slide_height / PT
    stamp = STAMP.search(prs.core_properties.subject or "")
    margin = MARGINS[960 if round(W) == 960 else 720].get(stamp.group(3) if stamp else "standard", 48)
    slides = []
    for number, slide in enumerate(prs.slides, 1):
        els = elements(slide, W, H)
        family = next((s.name.split("@", 1)[1] for s in slide.shapes if s.name.startswith("family@")), "")
        title = next((e for e in els if e["name"].startswith("title@")), None)
        s = {"n": number, "els": els, "family": family, "role": _role(family), "title": title,
             "named": title["name"].split("@", 1)[1] if title else None}
        s["kind"] = treatment(title, els, W, H) if title and title["sizes"] else None
        if s["role"] == "content":
            s["sig"] = signature(title, s["kind"], els, W, H, margin)
        slides.append(s)
    return slides


def assess(prs):
    findings = []

    def add(rule, severity, slide, value, threshold, note=""):
        findings.append({"rule": rule, "severity": severity, "slide": slide, "tier": "measured",
                         "value": value, "threshold": threshold, "note": note})

    W, H = prs.slide_width / PT, prs.slide_height / PT
    stamp = STAMP.search(prs.core_properties.subject or "")
    density, variance = (int(stamp.group(2)), int(stamp.group(4))) if stamp else (5, 5)
    slides = _read(prs)
    labels = [(n, s) for n, slide in enumerate(prs.slides, 1) for s in slide.shapes
              if s.name.startswith(("title@", "subtitle@")) and getattr(s, "has_text_frame", False) and s.has_text_frame]
    if not any(s["named"] for s in slides):
        return {"pass": True, "findings": [], "scope": "not drawn by the deck engine"}
    legacy = any(s["named"] == "legacy" for s in slides)
    variety = "MEDIUM" if legacy else "HIGH"

    for s in slides:
        if s["named"] in ZONES and s["kind"] not in (None, s["named"]):
            add("OF-110", "HIGH", s["n"], s["kind"], s["named"], "the title is drawn as one treatment and named as another")

    # OF-110: cover and section titles are variants of their own and do not count.
    titled = [s for s in slides if s["role"] not in ("cover", "section") and s["kind"]]
    used = sorted({s["kind"] for s in titled} - {"unclassified"})
    if len(slides) >= LIMIT["slides"]:
        if len(used) < LIMIT["treatments"]:
            add("OF-110", "MEDIUM" if variance <= 3 else variety, None, len(used), LIMIT["treatments"],
                "variance below 4 lowers variety on the user's request only" if variance <= 3 else ", ".join(used) or "none")
        elif stamp and len(used) > _most_treatments(variance):
            add("OF-110", variety, None, len(used), _most_treatments(variance), f"variance {variance}")
    lost = [s["n"] for s in titled if s["kind"] == "unclassified"]
    if lost:
        add("OF-110", variety, lost[0], len(lost), 0, "no treatment recognised on slide(s) " + ", ".join(map(str, lost)))

    # OF-111: the whole signature and its layout part (title zone x partition) both have to vary.
    content = [s for s in slides if s["role"] == "content"]
    if content:
        need = min(LIMIT["cap"], -(-len(content) * 3 // 5))
        for part, key in (("signature", lambda s: s["sig"]), ("layout", lambda s: s["sig"][:2])):
            counts = {}
            for s in content:
                counts[key(s)] = counts.get(key(s), 0) + 1
            top = max(counts.values())
            if len(counts) < need or top > LIMIT["share"] * len(content):
                add("OF-111", variety, None, {"distinct": len(counts), "largest_share": round(top / len(content), 2)},
                    {"distinct": need, "largest_share": LIMIT["share"]}, part)

    # OF-112: the band under the body, median over slides that have a body.
    cap = next(c for upto, c in BAND_CAPS if density <= upto)
    bands = []
    for s in slides:
        if s["role"] in ("cover", "section") or BAND_EXEMPT.fullmatch(s["family"] or "") or s["kind"] == "statement":
            continue
        value = band(s["title"], s["els"])
        if value is None:
            continue
        bands.append(value)
        if value > cap:
            add("OF-112", "MEDIUM", s["n"], round(value, 3), cap, "slide band over the density cap")
    if bands and statistics.median(bands) > LIMIT["band"]:
        add("OF-112", "MEDIUM" if legacy else "HIGH", None, round(statistics.median(bands), 3), LIMIT["band"], "deck median")

    # OF-113: one frame per treatment. A statement keeps its x and width (its height follows the lines); a
    # bottom title keeps its x, width and bottom edge, standing on the floor whatever its line count.
    frames = {}
    for s in slides:
        if s["kind"] and s["kind"] != "unclassified" and s["named"] not in ("cover", "section"):
            frames.setdefault(s["kind"], []).append(s)
    for kind, group in frames.items():
        first = group[0]["title"]
        keys = ("x", "w") if kind == "statement" else ("x", "w", "bottom") if kind == "bottom-anchor" else ("x", "y", "w")
        edge = lambda t, k: t["y"] + t["h"] if k == "bottom" else t[k]  # noqa: E731
        for s in group[1:]:
            drift = max(abs(edge(s["title"], k) - edge(first, k)) for k in keys)
            if drift > LIMIT["anchor"]:
                add("OF-113", "HIGH", s["n"], round(drift, 1), LIMIT["anchor"], f"{kind} frame moved against slide {group[0]['n']}")

    findings.extend(empty_regions(slides, W, H))
    findings.extend(light_figures(prs, W, H))
    findings.extend(run_ins(prs))
    findings.extend(agenda_numerals(slides))

    # OF-114: every engine-drawn title and the cover subtitle.
    for number, shape in labels:
        text = re.sub(r"\s+", " ", shape.text_frame.text).strip()
        if sentence(text):
            add("OF-114", "MEDIUM" if legacy else "HIGH", number, text[:60], "a topic label", shape.name)
    return {"pass": not any(f["severity"] == "HIGH" for f in findings), "findings": findings}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: deck_output.py deck.pptx")
    result = assess(Presentation(sys.argv[1]))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["pass"] else 1)
